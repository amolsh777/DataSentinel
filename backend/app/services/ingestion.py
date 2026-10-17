"""Safe parsing and validation of uploaded CSV / Excel files.

Uploaded content is only ever *parsed as data* by pandas/openpyxl; it is never
executed. Every rejection raises :class:`IngestionError` with a specific code
so the frontend can show an actionable message.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import PurePath

import pandas as pd

from app.core.errors import IngestionError

SUPPORTED_EXTENSIONS: dict[str, str] = {".csv": "csv", ".xlsx": "excel"}
_ZIP_MAGIC = b"PK\x03\x04"
_CANDIDATE_DELIMITERS = ",;\t|"
_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")
_NUMERIC_HEADER = re.compile(r"^-?\d+(\.\d+)?$")


@dataclass
class ParsedFile:
    """Result of parsing an uploaded file."""

    df: pd.DataFrame
    file_format: str
    sheet_name: str | None = None
    sheet_names: list[str] = field(default_factory=list)
    encoding: str | None = None
    delimiter: str | None = None
    warnings: list[str] = field(default_factory=list)


def detect_format(filename: str) -> str:
    """Map a filename's extension to a supported format or raise."""
    if not filename:
        raise IngestionError("A file name is required.", code="MISSING_FILENAME")
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise IngestionError(
            f"Unsupported file type '{ext or '(none)'}'. Supported: .csv, .xlsx.",
            code="UNSUPPORTED_FORMAT",
            details={"extension": ext, "supported": sorted(SUPPORTED_EXTENSIONS)},
        )
    return SUPPORTED_EXTENSIONS[ext]


def safe_display_filename(filename: str) -> str:
    """Strip any path components and control characters; metadata only, never a path."""
    name = PurePath(filename.replace("\\", "/")).name
    name = "".join(ch for ch in name if ch.isprintable())
    return name[:255] or "upload"


def parse_file(
    content: bytes,
    filename: str,
    *,
    sheet_name: str | None = None,
    max_rows: int = 1_000_000,
    max_columns: int = 500,
) -> ParsedFile:
    """Validate and parse an uploaded file into a DataFrame."""
    file_format = detect_format(filename)
    if not content or not content.strip():
        raise IngestionError("The file is empty.", code="EMPTY_FILE")
    if file_format == "csv":
        parsed = _parse_csv(content, max_rows=max_rows)
    else:
        parsed = _parse_excel(content, sheet_name=sheet_name, max_rows=max_rows)

    if parsed.df.shape[1] > max_columns:
        raise IngestionError(
            f"File has {parsed.df.shape[1]} columns; the limit is {max_columns}.",
            code="COLUMN_LIMIT_EXCEEDED",
            details={"columns": parsed.df.shape[1], "limit": max_columns},
        )
    return parsed


def list_excel_sheets(content: bytes) -> list[str]:
    """Return sheet names of an .xlsx workbook, raising on corrupted files."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    if not content.startswith(_ZIP_MAGIC):
        raise IngestionError(
            "The file is not a valid .xlsx workbook.", code="CORRUPTED_FILE"
        )
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        names = list(wb.sheetnames)
        wb.close()
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise IngestionError(
            "The Excel workbook could not be opened (corrupted or not .xlsx).",
            code="CORRUPTED_FILE",
            details={"reason": type(exc).__name__},
        ) from exc
    if not names:
        raise IngestionError("The workbook contains no sheets.", code="EMPTY_FILE")
    return names


# --------------------------------------------------------------------------- CSV


def _decode(content: bytes) -> tuple[str, str]:
    for enc in _ENCODINGS:
        try:
            return content.decode(enc), enc
        except UnicodeDecodeError:
            continue
    raise IngestionError("Could not decode the file text.", code="INVALID_CONTENT")  # pragma: no cover


def _sniff_delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:50])
    try:
        return csv.Sniffer().sniff(sample, delimiters=_CANDIDATE_DELIMITERS).delimiter
    except csv.Error:
        return ","


def _parse_csv(content: bytes, *, max_rows: int) -> ParsedFile:
    if content.startswith(_ZIP_MAGIC):
        raise IngestionError(
            "This looks like an Excel/zip file, not CSV. Upload it with a .xlsx extension.",
            code="INVALID_CONTENT",
        )
    if b"\x00" in content[:8192]:
        raise IngestionError(
            "The file appears to be binary, not text CSV. If it is UTF-16, re-save it as UTF-8.",
            code="INVALID_CONTENT",
        )
    text, encoding = _decode(content)
    warnings: list[str] = []
    if encoding != "utf-8-sig":
        warnings.append(f"File was not valid UTF-8; decoded as {encoding}.")

    delimiter = _sniff_delimiter(text)
    try:
        raw_header = next(csv.reader(io.StringIO(text), delimiter=delimiter))
    except StopIteration:
        raise IngestionError("The file is empty.", code="EMPTY_FILE") from None

    try:
        df = pd.read_csv(
            io.StringIO(text),
            sep=delimiter,
            header=0,
            low_memory=False,
            nrows=max_rows + 1,
        )
    except pd.errors.EmptyDataError as exc:
        raise IngestionError("The file has no columns.", code="EMPTY_FILE") from exc
    except pd.errors.ParserError as exc:
        raise IngestionError(
            "Malformed CSV: some rows have a different number of fields than the header.",
            code="MALFORMED_FILE",
            details={"parser_message": str(exc)[:300]},
        ) from exc

    if len(raw_header) != df.shape[1]:
        raise IngestionError(
            "Header and data column counts do not match.",
            code="MALFORMED_FILE",
            details={"header_columns": len(raw_header), "parsed_columns": df.shape[1]},
        )

    df = _apply_header(df, [str(h) for h in raw_header], warnings)
    _check_rows(df, max_rows)
    return ParsedFile(
        df=df, file_format="csv", encoding=encoding, delimiter=delimiter, warnings=warnings
    )


# ------------------------------------------------------------------------- Excel


def _parse_excel(content: bytes, *, sheet_name: str | None, max_rows: int) -> ParsedFile:
    sheets = list_excel_sheets(content)
    warnings: list[str] = []
    if sheet_name is None:
        sheet_name = sheets[0]
        if len(sheets) > 1:
            warnings.append(
                f"Workbook has {len(sheets)} sheets; using the first ('{sheet_name}'). "
                "Choose another sheet explicitly if needed."
            )
    elif sheet_name not in sheets:
        raise IngestionError(
            f"Sheet '{sheet_name}' not found.",
            code="SHEET_NOT_FOUND",
            details={"available_sheets": sheets},
        )

    try:
        raw = pd.read_excel(
            io.BytesIO(content),
            sheet_name=sheet_name,
            header=None,
            nrows=max_rows + 2,
            engine="openpyxl",
        )
    except Exception as exc:  # openpyxl raises a variety of types on bad content
        raise IngestionError(
            "The sheet could not be read.",
            code="CORRUPTED_FILE",
            details={"reason": type(exc).__name__},
        ) from exc

    raw = raw.dropna(how="all").reset_index(drop=True)  # Excel often has blank leading/trailing rows
    if raw.empty:
        raise IngestionError(f"Sheet '{sheet_name}' is empty.", code="EMPTY_DATASET")

    header = ["" if pd.isna(v) else str(v) for v in raw.iloc[0].tolist()]
    body = raw.iloc[1:].reset_index(drop=True).infer_objects()
    body.columns = range(body.shape[1])
    df = _apply_header(body, header, warnings)
    _check_rows(df, max_rows)
    return ParsedFile(
        df=df, file_format="excel", sheet_name=sheet_name, sheet_names=sheets, warnings=warnings
    )


# ------------------------------------------------------------------------ shared


def _apply_header(df: pd.DataFrame, raw_header: list[str], warnings: list[str]) -> pd.DataFrame:
    """Validate header names and assign them to ``df``.

    - Blank header + entirely empty column -> dropped with a warning (trailing delimiters).
    - Blank header + data -> MISSING_HEADER error.
    - Duplicate names (after trimming) -> DUPLICATE_COLUMNS error.
    """
    names = [h.strip() for h in raw_header]
    keep: list[int] = []
    missing_positions: list[int] = []
    for i, name in enumerate(names):
        if name:
            keep.append(i)
        elif df.iloc[:, i].isna().all():
            warnings.append(f"Dropped empty unnamed column at position {i + 1}.")
        else:
            missing_positions.append(i + 1)
    if missing_positions:
        raise IngestionError(
            "Some columns contain data but have no header name.",
            code="MISSING_HEADER",
            details={"column_positions": missing_positions},
        )

    df = df.iloc[:, keep].copy()
    kept_names = [names[i] for i in keep]
    if not kept_names:
        raise IngestionError("No named columns were found.", code="MISSING_HEADER")

    counts = Counter(kept_names)
    duplicates = sorted(n for n, c in counts.items() if c > 1)
    if duplicates:
        raise IngestionError(
            "Duplicate column names found. Rename them so each column is unique.",
            code="DUPLICATE_COLUMNS",
            details={"duplicate_columns": duplicates},
        )
    lowered = [n.lower() for n in kept_names]
    ci_dupes = sorted({n for n in kept_names if lowered.count(n.lower()) > 1})
    if ci_dupes:
        warnings.append(f"Column names differ only by case: {ci_dupes}.")
    if all(_NUMERIC_HEADER.match(n) for n in kept_names):
        warnings.append("All header values look numeric; the file may be missing a header row.")

    df.columns = kept_names
    return df


def _check_rows(df: pd.DataFrame, max_rows: int) -> None:
    if len(df) == 0:
        raise IngestionError("The file has a header but no data rows.", code="EMPTY_DATASET")
    if len(df) > max_rows:
        raise IngestionError(
            f"File has more than {max_rows:,} rows (the configured limit).",
            code="ROW_LIMIT_EXCEEDED",
            details={"limit": max_rows},
        )
